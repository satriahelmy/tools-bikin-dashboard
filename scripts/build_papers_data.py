"""Build the browser-facing Paper Library JSON from the editorial workbook.

This is a development-time script. The deployed site consumes only the generated
JSON file and never loads the workbook in the browser.
"""

from __future__ import annotations

import argparse
import json
import re
import sys
import unicodedata
from collections import Counter
from pathlib import Path
from typing import Any
from urllib.parse import urlparse

from openpyxl import load_workbook


DEFAULT_INPUT = Path("data/bikindashboard_paper_library_master_latest.xlsx")
DEFAULT_OUTPUT = Path("data/papers.json")
VERIFIED_ACCESS_STATUSES = {
    "arXiv",
    "Open Access",
    "Free PDF",
    "Verified Free PDF",
    "Free Full Text",
}
REJECTED_ACCESS_STATUSES = {
    "Metadata only",
    "No Free Full Text Verified",
}


def clean(value: Any) -> str | None:
    if value is None:
        return None
    value = str(value).strip()
    return value or None


def split_authors(value: Any) -> list[str]:
    return [author.strip() for author in (clean(value) or "").split(";") if author.strip()]


def normalize_importance(value: Any) -> str | None:
    value = clean(value)
    if value and value.lower() == "highly influential":
        return "Highly Influential"
    return value


def slugify(title: str) -> str:
    normalized = unicodedata.normalize("NFKD", title)
    normalized = normalized.encode("ascii", "ignore").decode("ascii")
    normalized = normalized.lower().replace("&", " and ")
    normalized = re.sub(r"[^a-z0-9]+", "-", normalized).strip("-")
    return normalized or "paper"


def is_http_url(value: str | None) -> bool:
    if not value:
        return False
    parsed = urlparse(value)
    return parsed.scheme in {"http", "https"} and bool(parsed.netloc)


def paper_is_verified(row: dict[str, Any]) -> bool:
    url = row["Verified Free Full-Text URL"]
    explicit_flag = (clean(row["Free Link Verified"]) or "").lower()
    access_status = clean(row["Access Status"])

    if not is_http_url(url):
        return False
    if explicit_flag == "no" or access_status in REJECTED_ACCESS_STATUSES:
        return False
    return explicit_flag == "yes" or access_status in VERIFIED_ACCESS_STATUSES


def build_record(row: dict[str, Any], used_ids: Counter[str]) -> dict[str, Any]:
    title = clean(row["Paper Title"])
    if not title:
        raise ValueError(f"Paper {row['No.']} is missing a title")

    base_id = slugify(title)
    used_ids[base_id] += 1
    paper_id = base_id if used_ids[base_id] == 1 else f"{base_id}-{used_ids[base_id]}"

    raw_paper_url = clean(row["Verified Free Full-Text URL"])
    paper_verified = paper_is_verified(row)
    raw_code_url = clean(row["Code / GitHub URL"])
    code_verified = bool(raw_code_url and clean(row["Code Verification"]) == "Verified")
    code_source = clean(row["Code Source"])
    if code_source and code_source.lower() == "none":
        code_source = None

    return {
        "id": paper_id,
        "source_no": int(row["No."]),
        "curated_order": int(row["No."]),
        "title": title,
        "year": int(row["Year"]),
        "authors": split_authors(row["Authors"]),
        "category": clean(row["Category"]),
        "topics": [],
        "contribution": clean(row["Main Contribution"]),
        "importance": normalize_importance(row["Importance"]),
        "difficulty": clean(row["Difficulty"]),
        "paper_url": raw_paper_url if paper_verified else None,
        "paper_source_type": clean(row["Free Source Type"]) or clean(row["Access Status"]),
        "paper_source_host": clean(row["Free Source Host"]),
        "paper_access_status": clean(row["Access Status"]),
        "paper_verified": paper_verified,
        "code_url": raw_code_url if code_verified else None,
        "code_source": code_source,
        "code_verified": code_verified,
    }


def read_rows(input_path: Path) -> list[dict[str, Any]]:
    workbook = load_workbook(input_path, read_only=True, data_only=True)
    if "Paper Catalog" not in workbook.sheetnames:
        raise ValueError("Workbook does not contain the required 'Paper Catalog' sheet")

    worksheet = workbook["Paper Catalog"]
    rows = worksheet.iter_rows(values_only=True)
    header = [clean(value) for value in next(rows)]
    if any(value is None for value in header):
        raise ValueError("Paper Catalog contains a blank header")

    records = []
    for values in rows:
        if not any(value is not None and str(value).strip() for value in values):
            continue
        records.append(dict(zip(header, values)))
    return records


def validate(records: list[dict[str, Any]], source_count: int) -> dict[str, Any]:
    required_fields = [
        "title",
        "year",
        "authors",
        "category",
        "contribution",
        "importance",
        "difficulty",
    ]
    missing = [
        record["source_no"]
        for record in records
        if any(not record[field] for field in required_fields)
    ]
    ids = [record["id"] for record in records]
    titles = [record["title"].casefold() for record in records]
    paper_urls = [record["paper_url"] for record in records if record["paper_url"]]
    code_urls = [record["code_url"] for record in records if record["code_url"]]
    duplicate_ids = [key for key, count in Counter(ids).items() if count > 1]
    duplicate_titles = [key for key, count in Counter(titles).items() if count > 1]
    duplicate_paper_urls = [key for key, count in Counter(paper_urls).items() if count > 1]
    duplicate_code_urls = [key for key, count in Counter(code_urls).items() if count > 1]
    invalid_urls = [
        record["source_no"]
        for record in records
        if (record["paper_url"] and not is_http_url(record["paper_url"]))
        or (record["code_url"] and not is_http_url(record["code_url"]))
    ]

    errors = {
        "source_count": source_count,
        "generated_count": len(records),
        "missing_required_fields": missing,
        "duplicate_ids": duplicate_ids,
        "duplicate_titles": duplicate_titles,
        "duplicate_paper_urls": duplicate_paper_urls,
        "duplicate_code_urls": duplicate_code_urls,
        "invalid_runtime_urls": invalid_urls,
    }
    if source_count != len(records) or any(errors[key] for key in errors if key not in {"source_count", "generated_count"}):
        raise ValueError(json.dumps(errors, ensure_ascii=False))

    return {
        **errors,
        "read_paper_count": sum(record["paper_verified"] for record in records),
        "code_count": sum(record["code_verified"] for record in records),
        "topic_value_count": sum(bool(record["topics"]) for record in records),
        "category_counts": dict(sorted(Counter(record["category"] for record in records).items())),
    }


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", type=Path, default=DEFAULT_INPUT)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()

    if not args.input.exists():
        raise FileNotFoundError(args.input)

    source_rows = read_rows(args.input)
    used_ids: Counter[str] = Counter()
    records = [build_record(row, used_ids) for row in source_rows]
    summary = validate(records, len(source_rows))

    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(
        json.dumps(records, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    print(json.dumps(summary, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except Exception as error:
        print(f"Paper data build failed: {error}", file=sys.stderr)
        raise
