"""Merge validated modern book candidates into the Book Library JSON dataset.

The workbook is an additions sheet, not a replacement for the existing curated
catalog. Existing records are preserved; only rows explicitly marked for
inclusion are appended in workbook order.
"""

from __future__ import annotations

import argparse
import json
from pathlib import Path
from urllib.parse import urlparse

import openpyxl


REQUIRED_HEADERS = {
    "Topic",
    "Title",
    "Author",
    "Level",
    "Access",
    "Format",
    "Official / Legal Link",
    "Include in Library?",
    "Why / Audit Note",
}


def clean(value: object) -> str:
    return str(value).strip() if value is not None else ""


def valid_url(value: str) -> bool:
    parsed = urlparse(value)
    return parsed.scheme in {"http", "https"} and bool(parsed.netloc)


def find_header(rows: list[tuple[object, ...]]) -> tuple[int, tuple[object, ...]]:
    for index, row in enumerate(rows):
        headers = {clean(value) for value in row if clean(value)}
        if REQUIRED_HEADERS.issubset(headers):
            return index, row
    raise ValueError("Could not find the validated book table header")


def read_additions(path: Path) -> list[dict[str, object]]:
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if "Modern AI Books" not in workbook.sheetnames:
        raise ValueError("Workbook does not contain the Modern AI Books sheet")

    rows = list(workbook["Modern AI Books"].iter_rows(values_only=True))
    header_index, header_row = find_header(rows)
    headers = [clean(value) for value in header_row]
    additions: list[dict[str, object]] = []

    for row_number, row in enumerate(rows[header_index + 1 :], start=header_index + 2):
        values = dict(zip(headers, row))
        title = clean(values.get("Title"))
        if not title:
            continue

        include_value = clean(values.get("Include in Library?")).lower()
        if not include_value.startswith("yes"):
            continue

        link = clean(values.get("Official / Legal Link"))
        if not valid_url(link):
            raise ValueError(f"Row {row_number} has an invalid Official / Legal Link: {link!r}")

        additions.append(
            {
                "title": title,
                "authors": clean(values.get("Author")),
                "category": clean(values.get("Topic")),
                "level": clean(values.get("Level")),
                "tools": "",
                "format": clean(values.get("Format")),
                "accessType": clean(values.get("Access")),
                "editionStatus": "",
                "verified": True,
                "url": link,
                "verificationUrl": link,
                "note": clean(values.get("Why / Audit Note")),
            }
        )

    return additions


def merge_catalog(source_path: Path, output_path: Path) -> dict[str, int]:
    payload = json.loads(output_path.read_text(encoding="utf-8"))
    books = payload.get("books")
    if not isinstance(books, list):
        raise ValueError("Existing Book Library JSON does not contain a books array")

    additions = read_additions(source_path)
    existing_titles = {clean(book.get("title")).casefold() for book in books}
    existing_urls = {
        clean(book.get("url")).rstrip("/").casefold()
        for book in books
        if clean(book.get("url"))
    }
    next_id = max((int(book.get("id", 0)) for book in books), default=0) + 1
    added = 0
    skipped_duplicates = 0

    for addition in additions:
        title_key = addition["title"].casefold()
        url_key = addition["url"].rstrip("/").casefold()
        if title_key in existing_titles or url_key in existing_urls:
            skipped_duplicates += 1
            continue
        books.append({"id": next_id, **addition})
        existing_titles.add(title_key)
        existing_urls.add(url_key)
        next_id += 1
        added += 1

    payload["meta"]["count"] = len(books)
    output_path.write_text(
        json.dumps(payload, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    return {
        "source_included": len(additions),
        "added": added,
        "skipped_duplicates": skipped_duplicates,
        "total": len(books),
    }


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--source",
        type=Path,
        default=Path("docs/source/book-library/workbooks/modern_ai_books_validated.xlsx"),
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=Path("data/books.json"),
    )
    args = parser.parse_args()
    print(json.dumps(merge_catalog(args.source, args.output), ensure_ascii=False))


if __name__ == "__main__":
    main()
